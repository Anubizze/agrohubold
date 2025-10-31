# users/models.py - обновленная модель с Word актом

from django.contrib.auth.models import AbstractUser
from django.db import models
from django.core.files.base import ContentFile
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from django.conf import settings
import io
import os
from docx.oxml.shared import qn
from docx.oxml import OxmlElement
class User(AbstractUser):
	phone = models.CharField(max_length=20, blank=True, verbose_name="Телефон")
	
	def __str__(self):
		return f"{self.get_full_name()} ({self.username})"
	
	class Meta:
		verbose_name = "Пользователь"
		verbose_name_plural = "Пользователи"


class ServiceCart(models.Model):
	"""Корзина услуг пользователя"""
	user = models.ForeignKey(User, on_delete=models.CASCADE, verbose_name="Пользователь")
	service = models.ForeignKey('main_app.Service', on_delete=models.CASCADE, verbose_name="Услуга")
	added_at = models.DateTimeField(auto_now_add=True, verbose_name="Добавлено")
	
	def __str__(self):
		return f"{self.user.username} - {self.service.name}"
	
	class Meta:
		verbose_name = "Корзина услуг"
		verbose_name_plural = "Корзина услуг"
		unique_together = ('user', 'service')
		ordering = ['-added_at']


class ServicePayment(models.Model):
	"""Платежи за услуги"""
	STATUS_CHOICES = [
		('pending', 'Ожидает проверки'),
		('approved', 'Оплачено'),
		('rejected', 'Отклонено'),
	]
	
	user = models.ForeignKey(User, on_delete=models.CASCADE, verbose_name="Пользователь")
	services = models.ManyToManyField('main_app.Service', verbose_name="Услуги")
	
	# Информация об оплате
	total_amount = models.DecimalField(max_digits=10, decimal_places=2, verbose_name="Сумма")
	receipt_file = models.FileField(upload_to='receipts/', verbose_name="Чек")
	
	# Акт выполненных работ (Word файл)
	act_file = models.FileField(upload_to='acts/', blank=True, null=True, verbose_name="Акт выполненных работ")
	
	# Статус и даты
	status = models.CharField(max_length=20, choices=STATUS_CHOICES, default='pending', verbose_name="Статус")
	created_at = models.DateTimeField(auto_now_add=True, verbose_name="Создан")
	updated_at = models.DateTimeField(auto_now=True, verbose_name="Обновлен")
	
	def __str__(self):
		services_count = self.services.count()
		return f"Платеж {self.id} - {self.user.username} ({services_count} услуг) - {self.total_amount}₸"
	
	def get_services_list(self):
		"""Возвращает список названий услуг для админки"""
		return ", ".join([service.name for service in self.services.all()])
	get_services_list.short_description = 'Услуги'
	
	def get_act_number(self):
		"""Генерирует номер акта: дата + ID платежа"""
		return f"{self.created_at.strftime('%d%m%Y')}{self.id:06d}"
	
	def get_providers_list(self):
		"""Возвращает список руководителей (поставщиков услуг)"""
		providers = set()
		for service in self.services.all():
			provider = getattr(service.category, 'provider', None)
			if provider:
				providers.add(provider.name)
		return ", ".join(sorted(providers)) or "—"
	get_providers_list.short_description = "Поставщики"
	

	def generate_act_file(self):
		"""Генерирует Word файл акта"""
		doc = Document()
		style = doc.styles['Normal']
		font = style.font
		font.name = 'Times New Roman'

		act_number = self.get_act_number()
		title = doc.add_paragraph(f'АКТ ВЫПОЛНЕННЫХ РАБОТ (ОКАЗАННЫХ УСЛУГ) № {act_number}')
		_style_heading(title)

		date_time = doc.add_paragraph()
		date_time.add_run(f'Дата составления акта: {self.created_at.strftime("%d.%m.%Y")}').bold = True

		# --- Заказчик ---
		customer_table = doc.add_table(rows=1, cols=2)
		customer_table.cell(0, 0).text = 'Заказчик:'
		user_name = self.user.get_full_name() or self.user.username
		customer_info = f'''Имя: {user_name}
Email: {self.user.email or "Не указан"}
Телефон: {self.user.phone or "Не указан"}'''
		customer_table.cell(0, 1).text = customer_info
		doc.add_paragraph()

		# --- Исполнитель ---
		# Получаем список уникальных поставщиков
		providers = list(
			set(service.category.provider for service in self.services.all())
		)
		providers_names = ", ".join(p.name for p in providers)
		executor_text = f'''Некоммерческое акционерное общество «Шәкәрім университет»
Поставщик(и): {providers_names}'''

		executor_table = doc.add_table(rows=1, cols=2)
		executor_table.cell(0, 0).text = 'Исполнитель:'
		executor_table.cell(0, 1).text = executor_text
		doc.add_paragraph()

		# --- Таблица услуг ---
		services_title = doc.add_paragraph('Перечень выполненных работ (оказанных услуг)')
		_style_heading(services_title)

		services_table = doc.add_table(rows=1, cols=7)
		services_table.style = 'Table Grid'

		headers = ['№', 'Наименование услуги', 'Оператор', 'Ед. изм.', 'Кол-во', 'Цена, тг', 'Сумма, тг']
		for i, header in enumerate(headers):
			cell = services_table.cell(0, i)
			cell.text = header
			cell.paragraphs[0].runs[0].bold = True

		for i, service in enumerate(self.services.all(), 1):
			row = services_table.add_row()
			operator_name = service.operator.get_full_name() or service.operator.username
			row.cells[0].text = str(i)
			row.cells[1].text = service.name
			row.cells[2].text = operator_name
			row.cells[3].text = 'услуга'
			row.cells[4].text = '1'
			row.cells[5].text = f'{service.price:,.0f}'
			row.cells[6].text = f'{service.price:,.2f}'

		doc.add_paragraph()
		total_p = doc.add_paragraph()
		total_p.add_run(f'Итого к оплате: {self.total_amount:,.2f} тг').bold = True

		doc.add_paragraph()
		paid_p = doc.add_paragraph()
		paid_p.add_run('Оплачено через Kaspi QR: ✔').bold = True

		domain = getattr(settings, 'SITE_DOMAIN')
		receipt_url = f"\nЧек: {domain}{self.receipt_file.url}"
		paid_p.add_run(f'{receipt_url}')

		doc.add_paragraph()
		doc.add_paragraph(f'Документ сформирован автоматически системой учета услуг сайта Agrosector ({domain})')
		doc.add_paragraph('НАО «Шәкәрім университет».')
		
		# --- Сохранение ---
		doc_io = io.BytesIO()
		doc.save(doc_io)
		doc_io.seek(0)
		filename = f'act_{act_number}.docx'
		self.act_file.save(filename, ContentFile(doc_io.read()), save=False)

		return self.act_file
	
	class Meta:
		verbose_name = "Платеж за услуги"
		verbose_name_plural = "Платежи за услуги"
		ordering = ['-created_at']

def _style_heading(heading_paragraph):
	"""Применяет единый стиль к заголовку"""
	heading_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
	for run in heading_paragraph.runs:
		run.font.bold = True
		run.font.name = 'Times New Roman'